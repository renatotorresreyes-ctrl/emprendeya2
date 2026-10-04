"""Genera las 8 plantillas .xlsx simples (sin macros).
Uso: python3 scripts/make_templates.py
Luego recalcula cada archivo para que traiga valores en caché."""
import datetime
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(__file__), "..", "public", "plantillas")
F = "Arial"
HEAD = PatternFill("solid", fgColor="14213D")
INP = PatternFill("solid", fgColor="FFF9DB")
thin = Side(style="thin", color="DCE3EC")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CRC = '"₡"#,##0;("₡"#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
N = 60  # filas disponibles para llenar
LEG = "Celdas amarillas = las llenas tú. Celdas blancas = fórmulas (no las borres). La fila 5 es un ejemplo: reemplázala o bórrala."


def base(title, note, widths):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plantilla"
    ws["A1"] = title
    ws["A1"].font = Font(name=F, size=14, bold=True, color="14213D")
    ws["A2"] = note
    ws["A2"].font = Font(name=F, size=10, italic=True, color="55627A")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    return wb, ws


def header(ws, row, names):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30


def cell(ws, ref, value=None, fmt=None, inp=False, bold=False, blue=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name=F, size=10, bold=bold, color="0000FF" if blue else "000000")
    if fmt:
        c.number_format = fmt
    if inp:
        c.fill = INP
    c.border = BOX
    return c


def note(ws, ref, text):
    ws[ref] = text
    ws[ref].font = Font(name=F, size=9, italic=True, color="55627A")


def label(ws, ref, text, bold=True):
    ws[ref] = text
    ws[ref].font = Font(name=F, size=10, bold=bold)


def save(wb, name):
    wb.save(os.path.join(OUT, name + ".xlsx"))


# 1 Registro de ventas
wb, ws = base("Registro de ventas", LEG, [14, 26, 12, 20, 18, 4, 20])
header(ws, 4, ["Fecha", "Producto", "Cantidad", "Precio unitario (₡)", "Total (₡)"])
for r in range(5, 5 + N):
    cell(ws, f"A{r}", fmt="dd/mm/yyyy", inp=True)
    cell(ws, f"B{r}", inp=True)
    cell(ws, f"C{r}", fmt="#,##0", inp=True)
    cell(ws, f"D{r}", fmt=CRC, inp=True)
    cell(ws, f"E{r}", f'=IF(OR(C{r}="",D{r}=""),"",C{r}*D{r})', CRC)
ws["A5"] = datetime.date(2026, 1, 5)
ws["B5"] = "Ejemplo: producto"
ws["C5"] = 2
ws["D5"] = 3500
label(ws, "G4", "Total vendido (₡)")
cell(ws, "G5", f"=SUM(E5:E{4 + N})", CRC, bold=True)
ws.freeze_panes = "A5"
save(wb, "registro-de-ventas")

# 2 Control de gastos
wb, ws = base("Control de gastos", LEG, [14, 26, 34, 16, 24, 4, 20])
header(ws, 4, ["Fecha", "Categoría", "Descripción", "Monto (₡)", "Método de pago"])
dv1 = DataValidation(
    type="list",
    formula1='"Materia prima,Alquiler,Servicios públicos,Transporte,Marketing,Comisiones,Salarios,Impuestos,Mantenimiento,Software,Otros"',
    allow_blank=True,
)
dv2 = DataValidation(type="list", formula1='"Efectivo,SINPE / Transferencia,Tarjeta,Otro"', allow_blank=True)
ws.add_data_validation(dv1)
ws.add_data_validation(dv2)
for r in range(5, 5 + N):
    cell(ws, f"A{r}", fmt="dd/mm/yyyy", inp=True)
    cell(ws, f"B{r}", inp=True)
    cell(ws, f"C{r}", inp=True)
    cell(ws, f"D{r}", fmt=CRC, inp=True)
    cell(ws, f"E{r}", inp=True)
    dv1.add(f"B{r}")
    dv2.add(f"E{r}")
ws["A5"] = datetime.date(2026, 1, 3)
ws["B5"] = "Alquiler"
ws["C5"] = "Ejemplo: alquiler del local"
ws["D5"] = 150000
ws["E5"] = "SINPE / Transferencia"
label(ws, "G4", "Total gastado (₡)")
cell(ws, "G5", f"=SUM(D5:D{4 + N})", CRC, bold=True)
ws.freeze_panes = "A5"
save(wb, "control-de-gastos")

# 3 Costos por producto
wb, ws = base("Costos por producto", LEG, [28, 20, 20, 18, 14])
header(ws, 4, ["Producto", "Costo variable (₡)", "Precio de venta (₡)", "Margen (₡)", "Margen (%)"])
for r in range(5, 5 + N):
    cell(ws, f"A{r}", inp=True)
    cell(ws, f"B{r}", fmt=CRC, inp=True)
    cell(ws, f"C{r}", fmt=CRC, inp=True)
    cell(ws, f"D{r}", f'=IF(OR(B{r}="",C{r}=""),"",C{r}-B{r})', CRC)
    cell(ws, f"E{r}", f'=IF(OR(B{r}="",C{r}="",C{r}=0),"",(C{r}-B{r})/C{r})', PCT)
ws["A5"] = "Ejemplo: producto"
ws["B5"] = 1200
ws["C5"] = 3500
ws.freeze_panes = "A5"
save(wb, "costos-por-producto")

# 4 Punto de equilibrio
wb, ws = base("Punto de equilibrio", "Llena las 3 celdas amarillas (los valores son un ejemplo). El resultado se calcula solo.", [40, 22])
for i, (lab, val) in enumerate(
    [("Precio de venta por unidad (₡)", 3500), ("Costo variable por unidad (₡)", 1200), ("Costos fijos del mes (₡)", 1800000)], 4
):
    label(ws, f"A{i}", lab, bold=False)
    cell(ws, f"B{i}", val, CRC, inp=True, blue=True)
rows = [
    ("Margen de contribución por unidad (₡)", "=B4-B5", CRC),
    ("Unidades de equilibrio por mes", '=IF(B4>B5,ROUNDUP(B6/(B4-B5),0),"Precio debe ser mayor al costo")', "#,##0"),
    ("Ventas de equilibrio (₡)", '=IF(B4>B5,B9*B4,"")', CRC),
]
for i, (lab, f, fmt) in enumerate(rows, 8):
    label(ws, f"A{i}", lab)
    cell(ws, f"B{i}", f, fmt, bold=True)
note(ws, "A12", "Fórmula: unidades = costos fijos / (precio − costo variable).")
save(wb, "punto-de-equilibrio")

# 5 Flujo de caja
wb, ws = base(
    "Flujo de caja mensual",
    "Escribe el saldo inicial de enero y los ingresos y egresos de cada mes. Cada mes arranca con el saldo final del anterior.",
    [14, 20, 18, 18, 20],
)
header(ws, 4, ["Mes", "Saldo inicial (₡)", "Ingresos (₡)", "Egresos (₡)", "Saldo final (₡)"])
meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
for i, m in enumerate(meses):
    r = 5 + i
    cell(ws, f"A{r}", m, bold=True)
    cell(ws, f"B{r}", 0 if i == 0 else f"=E{r - 1}", CRC, inp=(i == 0))
    cell(ws, f"C{r}", 0, CRC, inp=True)
    cell(ws, f"D{r}", 0, CRC, inp=True)
    cell(ws, f"E{r}", f"=B{r}+C{r}-D{r}", CRC)
note(ws, "A18", "Si el saldo final sale negativo, ese mes necesitarás financiamiento o recortar gastos.")
save(wb, "flujo-de-caja")

# 6 Inventario
wb, ws = base("Control de inventario", LEG, [12, 26, 14, 12, 20, 18, 14, 4, 26])
header(ws, 4, ["Código", "Producto", "Existencias", "Mínimo", "Costo unitario (₡)", "Valor total (₡)", "Estado"])
for r in range(5, 5 + N):
    cell(ws, f"A{r}", inp=True)
    cell(ws, f"B{r}", inp=True)
    cell(ws, f"C{r}", fmt="#,##0", inp=True)
    cell(ws, f"D{r}", fmt="#,##0", inp=True)
    cell(ws, f"E{r}", fmt=CRC, inp=True)
    cell(ws, f"F{r}", f'=IF(OR(C{r}="",E{r}=""),"",C{r}*E{r})', CRC)
    cell(ws, f"G{r}", f'=IF(OR(C{r}="",D{r}=""),"",IF(C{r}<=D{r},"Reponer","OK"))')
ws["A5"] = "001"
ws["B5"] = "Ejemplo: producto"
ws["C5"] = 10
ws["D5"] = 3
ws["E5"] = 500
label(ws, "I4", "Valor del inventario (₡)")
cell(ws, "I5", f"=SUM(F5:F{4 + N})", CRC, bold=True)
ws.freeze_panes = "A5"
save(wb, "inventario")

# 7 Presupuesto mensual
wb, ws = base("Presupuesto mensual", "Celdas amarillas = las llenas tú. Escribe lo presupuestado y lo que realmente gastaste.", [30, 20, 18, 18, 14])
header(ws, 4, ["Rubro", "Presupuestado (₡)", "Real (₡)", "Diferencia (₡)", "% usado"])
rub = ["Alquiler", "Materia prima / Mercadería", "Servicios públicos", "Transporte", "Marketing", "Salarios", "Impuestos", "Otros"]
for i, nme in enumerate(rub):
    r = 5 + i
    cell(ws, f"A{r}", nme, inp=True)
    cell(ws, f"B{r}", 0, CRC, inp=True)
    cell(ws, f"C{r}", 0, CRC, inp=True)
    cell(ws, f"D{r}", f"=B{r}-C{r}", CRC)
    cell(ws, f"E{r}", f"=IF(B{r}=0,0,C{r}/B{r})", PCT)
t = 5 + len(rub)
cell(ws, f"A{t}", "TOTAL", bold=True)
for col in "BCD":
    cell(ws, f"{col}{t}", f"=SUM({col}5:{col}{t - 1})", CRC, bold=True)
cell(ws, f"E{t}", f"=IF(B{t}=0,0,C{t}/B{t})", PCT, bold=True)
note(ws, f"A{t + 2}", "Diferencia positiva = ahorraste. Negativa = te pasaste del presupuesto.")
save(wb, "presupuesto-mensual")

# 8 Plan financiero a 12 meses
wb, ws = base("Plan financiero a 12 meses", "Proyecta ventas, costos y gastos de cada mes (celdas amarillas). La utilidad se calcula sola.", [10, 18, 18, 18, 18, 24])
header(ws, 4, ["Mes", "Ventas (₡)", "Costos (₡)", "Gastos (₡)", "Utilidad (₡)", "Utilidad acumulada (₡)"])
for i in range(12):
    r = 5 + i
    cell(ws, f"A{r}", i + 1, bold=True)
    for col in "BCD":
        cell(ws, f"{col}{r}", 0, CRC, inp=True)
    cell(ws, f"E{r}", f"=B{r}-C{r}-D{r}", CRC)
    cell(ws, f"F{r}", f"=E{r}" if i == 0 else f"=F{r - 1}+E{r}", CRC)
cell(ws, "A17", "TOTAL", bold=True)
for col in "BCDE":
    cell(ws, f"{col}17", f"=SUM({col}5:{col}16)", CRC, bold=True)
note(ws, "A19", "Costos = lo que cuesta producir lo que vendes. Gastos = alquiler, servicios, salarios, etc.")
save(wb, "plan-financiero-12-meses")
print("ok")
